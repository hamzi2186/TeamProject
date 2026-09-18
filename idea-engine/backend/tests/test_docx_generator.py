import os
import tempfile
import uuid
from datetime import date, datetime, timezone

import docx

from app.schemas.enums import Channel, LeadOutcome, ReportStatus
from app.schemas.reports import (
    IdeaReportRunResponse,
    IdeaReportSummaryCounts,
    LeadReportItemResponse,
    TimelineEvent,
)
from app.services.docx_generator import build_daily_docx_report


def test_docx_generation():
    today = date(2026, 9, 16)
    now = datetime(2026, 9, 16, 17, 0, tzinfo=timezone.utc)

    item = LeadReportItemResponse(
        id=uuid.uuid4(),
        lead_id=uuid.uuid4(),
        lead_name="Jane Doe (Acme Corp)",
        company_website="https://acme.example.com",
        phone="+15551234567",
        email="jane@acme.example.com",
        final_outcome=LeadOutcome.INTERESTED,
        approach_summary="First reached via Email, followed by SMS.",
        conversation_summary="Lead expressed excitement about the autonomous outreach capabilities.",
        outcome_reason="Lead requested demo link and enterprise pricing schedule.",
        recommended_next_action="Deliver demo link and schedule discovery session.",
        first_activity_at=now,
        last_activity_at=now,
        source_event_count=2,
        channels_used=[Channel.EMAIL, Channel.SMS],
        campaigns=["Q3 Enterprise Outreach"],
        timeline=[],
    )

    counts = IdeaReportSummaryCounts(
        total_leads=1,
        interested=1,
        converted=0,
        not_interested=0,
        follow_up_required=0,
        no_answer=0,
        no_response=0,
        do_not_contact=0,
        failed=0,
    )

    report_run = IdeaReportRunResponse(
        id=uuid.uuid4(),
        report_date=today,
        status=ReportStatus.COMPLETED,
        total_leads=1,
        summary_counts=counts,
        generated_at=now,
        document_filename="test-report.docx",
        download_url="/idea/reports/test/download",
        created_at=now,
        updated_at=now,
        items=[item],
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        output_file = os.path.join(tmpdir, "test_output.docx")
        result_path = build_daily_docx_report(report_run, output_file)

        assert os.path.exists(result_path)
        assert os.path.getsize(result_path) > 1000

        # Read back with docx library
        doc = docx.Document(result_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        combined = text + "\n" + table_text
        assert "T REX — DAILY LEAD INTELLIGENCE REPORT" in combined
        assert "Jane Doe (Acme Corp)" in combined
        assert "INTERESTED" in combined
        assert "Q3 Enterprise Outreach" in table_text
        assert "Factual Event Log" in text or "No persisted communication events found" in text
        assert len(doc.tables) >= 2
