import os
from datetime import date, datetime, timezone
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.reports import IdeaReportRunResponse, LeadReportItemResponse, TimelineEvent


def _set_cell_background(cell, color_hex: str) -> None:
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def _style_cell_text(cell, size=9.5, color=(0x11, 0x13, 0x0F), bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(*color)
            r.font.bold = bold


def build_daily_docx_report(
    report_run: IdeaReportRunResponse,
    output_path: str | Path,
) -> str:
    """
    Constructs a highly polished, professional .docx lead intelligence report
    adhering to the T Rex design aesthetic and PRD Section 2.6 requirements.

    Layout:
    - Page 1: document title, report meta block
    - Page 2: executive summary / outcome KPI breakdown
    - One dedicated page per lead: identity, campaigns, channels, factual
      event log, generated narratives, outcome evidence, next action.
    """
    doc = docx.Document()

    # Set 0.75-inch page margins for professional appearance
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Document Title / Header
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("T REX — DAILY LEAD INTELLIGENCE REPORT")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x11, 0x13, 0x0F)
    title_p.paragraph_format.space_after = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Autonomous Multi-Channel Outreach & Lead Conversion Summary")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x4F, 0x52, 0x4A)
    subtitle_p.paragraph_format.space_after = Pt(14)

    # Meta Info block
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    gen_time_str = (
        report_run.generated_at.strftime("%Y-%m-%d %H:%M:%S UTC")
        if report_run.generated_at
        else datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    )

    cells = meta_table.rows[0].cells
    cells[0].text = f"Report Date: {report_run.report_date.strftime('%B %d, %Y')}"
    cells[1].text = f"Generated: {gen_time_str}"
    cells2 = meta_table.rows[1].cells
    cells2[0].text = f"Total Leads Reviewed: {report_run.total_leads}"
    cells2[1].text = f"Report Status: {report_run.status.value}"

    for row in meta_table.rows:
        for cell in row.cells:
            _set_cell_background(cell, "F1EFE9")
            _set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            _style_cell_text(cell)

    # Start the executive summary on its own page
    doc.add_page_break()

    # 2. Executive Summary / KPI Breakdown
    kpi_heading = doc.add_paragraph()
    kh_run = kpi_heading.add_run("1. Executive Summary & Outcome Breakdown")
    kh_run.font.name = "Arial"
    kh_run.font.size = Pt(14)
    kh_run.font.bold = True
    kh_run.font.color.rgb = RGBColor(0x1F, 0x5A, 0x3A)
    kpi_heading.paragraph_format.space_after = Pt(6)

    kpi_table = doc.add_table(rows=1, cols=3)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = kpi_table.rows[0].cells
    hdr_cells[0].text = "Outcome Category"
    hdr_cells[1].text = "Lead Count"
    hdr_cells[2].text = "% of Reviewed Leads"

    for cell in hdr_cells:
        _set_cell_background(cell, "1F5A3A")
        _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    counts = report_run.summary_counts
    total = max(report_run.total_leads, 1)

    kpi_rows_data = [
        ("Interested", counts.interested, f"{(counts.interested / total) * 100:.1f}%"),
        ("Converted", counts.converted, f"{(counts.converted / total) * 100:.1f}%"),
        ("Follow-up Required", counts.follow_up_required, f"{(counts.follow_up_required / total) * 100:.1f}%"),
        ("Not Interested", counts.not_interested, f"{(counts.not_interested / total) * 100:.1f}%"),
        ("No Answer", counts.no_answer, f"{(counts.no_answer / total) * 100:.1f}%"),
        ("No Response", counts.no_response, f"{(counts.no_response / total) * 100:.1f}%"),
        ("Do Not Contact", counts.do_not_contact, f"{(counts.do_not_contact / total) * 100:.1f}%"),
        ("Failed / Other", counts.failed + counts.new + counts.contacting, f"{((counts.failed + counts.new + counts.contacting) / total) * 100:.1f}%"),
    ]

    for idx, (label, count, pct) in enumerate(kpi_rows_data):
        row_cells = kpi_table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = str(count)
        row_cells[2].text = pct
        bg_color = "FFFFFF" if idx % 2 == 0 else "F7F5F0"
        for cell in row_cells:
            _set_cell_background(cell, bg_color)
            _set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            _style_cell_text(cell)

    # 3. Individual Lead Dossiers (one page per lead)
    doc.add_page_break()

    leads_heading = doc.add_paragraph()
    lh_run = leads_heading.add_run("2. Individual Lead Intelligence Dossiers (One Page Per Lead)")
    lh_run.font.name = "Arial"
    lh_run.font.size = Pt(14)
    lh_run.font.bold = True
    lh_run.font.color.rgb = RGBColor(0x1F, 0x5A, 0x3A)
    leads_heading.paragraph_format.space_after = Pt(8)

    if not report_run.items:
        no_leads_p = doc.add_paragraph()
        no_leads_p.add_run("No lead interactions recorded for this date.")
    else:
        doc.add_page_break()
        for idx, item in enumerate(report_run.items, start=1):
            _add_lead_dossier_section(doc, idx, item)
            if idx != len(report_run.items):
                doc.add_page_break()

    # Save to disk
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(str(output_path))
    return str(output_path)


def _add_lead_dossier_section(doc: docx.Document, index: int, item: LeadReportItemResponse) -> None:
    # Lead Banner
    lead_title_p = doc.add_paragraph()
    lead_title_p.paragraph_format.space_before = Pt(4)
    lead_title_p.paragraph_format.space_after = Pt(4)

    r_num = lead_title_p.add_run(f"Lead {index}: {item.lead_name} ")
    r_num.font.name = "Arial"
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x11, 0x13, 0x0F)

    r_outcome = lead_title_p.add_run(f"[{item.final_outcome.value}]")
    r_outcome.font.name = "Arial"
    r_outcome.font.size = Pt(11)
    r_outcome.font.bold = True
    r_outcome.font.color.rgb = _get_outcome_color(item.final_outcome.value)

    # Lead Contact Summary Table
    info_table = doc.add_table(rows=3, cols=3)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    r1 = info_table.rows[0].cells
    r1[0].text = f"Email: {item.email or 'N/A'}"
    r1[1].text = f"Phone: {item.phone or 'N/A'}"
    r1[2].text = f"Website: {item.company_website or 'N/A'}"

    campaigns_str = ", ".join(item.campaigns) if item.campaigns else "None"
    channels_str = ", ".join(c.value for c in item.channels_used) if item.channels_used else "None"
    r2 = info_table.rows[1].cells
    r2[0].text = f"Campaigns: {campaigns_str}"
    r2[1].text = f"Channels: {channels_str}"
    r2[2].text = f"Total Touchpoints: {item.source_event_count}"

    first_touch_str = item.first_activity_at.strftime("%m/%d %H:%M") if item.first_activity_at else "N/A"
    last_touch_str = item.last_activity_at.strftime("%m/%d %H:%M") if item.last_activity_at else "N/A"
    r3 = info_table.rows[2].cells
    r3[0].text = f"First Touch: {first_touch_str}"
    r3[1].text = f"Last Touch: {last_touch_str}"
    r3[2].text = f"Current Status: {item.final_outcome.value}"

    for row in info_table.rows:
        for cell in row.cells:
            _set_cell_background(cell, "F7F5F0")
            _set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            _style_cell_text(cell, size=8.5, color=(0x4F, 0x52, 0x4A))

    # Factual Event Log (Persisted Records) - clear separation from generated summary
    log_heading = doc.add_paragraph()
    log_heading.paragraph_format.space_before = Pt(10)
    log_heading.paragraph_format.space_after = Pt(4)
    lh_run = log_heading.add_run("Factual Event Log (Persisted Records)")
    lh_run.font.name = "Arial"
    lh_run.font.size = Pt(10.5)
    lh_run.font.bold = True
    lh_run.font.color.rgb = RGBColor(0x1F, 0x5A, 0x3A)
    _add_event_log_table(doc, item.timeline)

    # Approach Narrative (generated)
    p_approach = doc.add_paragraph()
    p_approach.paragraph_format.space_before = Pt(8)
    p_approach.paragraph_format.space_after = Pt(2)
    lbl_app = p_approach.add_run("Approach Narrative (Generated): ")
    lbl_app.font.bold = True
    lbl_app.font.size = Pt(9.5)
    body_app = p_approach.add_run(item.approach_summary or "Outreach not initiated.")
    body_app.font.size = Pt(9.5)

    # Conversation Synthesis (generated)
    p_conv = doc.add_paragraph()
    p_conv.paragraph_format.space_before = Pt(4)
    p_conv.paragraph_format.space_after = Pt(2)
    lbl_conv = p_conv.add_run("Conversation Synthesis (Generated): ")
    lbl_conv.font.bold = True
    lbl_conv.font.size = Pt(9.5)
    body_conv = p_conv.add_run(item.conversation_summary or "No active conversation turns.")
    body_conv.font.size = Pt(9.5)

    # Outcome Reason / Evidence
    p_reason = doc.add_paragraph()
    p_reason.paragraph_format.space_before = Pt(4)
    p_reason.paragraph_format.space_after = Pt(2)
    lbl_rsn = p_reason.add_run("Outcome Justification & Evidence: ")
    lbl_rsn.font.bold = True
    lbl_rsn.font.size = Pt(9.5)
    body_rsn = p_reason.add_run(item.outcome_reason or "Assigned based on default campaign rules.")
    body_rsn.font.size = Pt(9.5)

    # Recommended Next Action Box
    action_table = doc.add_table(rows=1, cols=1)
    action_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = action_table.rows[0].cells[0]
    _set_cell_background(c, "F1EFE9")
    _set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    p_act = c.paragraphs[0]
    r_act_lbl = p_act.add_run("Recommended Next Action: ")
    r_act_lbl.font.bold = True
    r_act_lbl.font.size = Pt(9.5)
    r_act_lbl.font.color.rgb = RGBColor(0x1F, 0x5A, 0x3A)
    r_act_val = p_act.add_run(item.recommended_next_action or "Review lead status.")
    r_act_val.font.size = Pt(9.5)
    r_act_val.font.italic = True


def _add_event_log_table(doc: docx.Document, events: list[TimelineEvent]) -> None:
    if not events:
        p = doc.add_paragraph()
        run = p.add_run("No persisted communication events found for this lead.")
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x68, 0x6A, 0x63)
        return

    headers = ["Time (UTC)", "Channel", "Dir", "From", "To", "Subject", "Content", "Details"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        _set_cell_background(cell, "1F5A3A")
        _set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, event in enumerate(events):
        ts = event.timestamp.strftime("%Y-%m-%d %H:%M:%S") if event.timestamp else "N/A"
        content = (event.content or "").replace("\n", " ").strip()
        if len(content) > 120:
            content = content[:117] + "..."
        subject = (event.subject or "")[:40] if event.subject else ""
        details_bits = []
        if event.duration_seconds is not None:
            details_bits.append(f"{event.duration_seconds}s")
        if event.delivery_status:
            details_bits.append(event.delivery_status)
        if event.outcome:
            details_bits.append(event.outcome)

        row_values = [
            ts,
            event.channel.value if event.channel else "N/A",
            event.direction.value if event.direction else "N/A",
            event.sender or "-",
            event.recipient or "-",
            subject or "-",
            content or "-",
            ", ".join(details_bits) or "-",
        ]
        row = table.add_row().cells
        for idx, value in enumerate(row_values):
            row[idx].text = value
            _set_cell_background(row[idx], "FFFFFF" if i % 2 == 0 else "F7F5F0")
            _set_cell_margins(row[idx], top=30, bottom=30, left=60, right=60)
            for p in row[idx].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)


def _get_outcome_color(outcome: str) -> RGBColor:
    match outcome.upper():
        case "INTERESTED" | "CONVERTED":
            return RGBColor(0x1F, 0x6A, 0x45)  # Success green
        case "FOLLOW_UP_REQUIRED" | "CONTACTING" | "SCHEDULED":
            return RGBColor(0xC9, 0x89, 0x24)  # Warning amber
        case "DO_NOT_CONTACT" | "FAILED":
            return RGBColor(0xC9, 0x3A, 0x32)  # Danger red
        case "NOT_INTERESTED" | "NO_ANSWER" | "NO_RESPONSE":
            return RGBColor(0x68, 0x6A, 0x63)  # Muted grey
        case _:
            return RGBColor(0x4F, 0x52, 0x4A)